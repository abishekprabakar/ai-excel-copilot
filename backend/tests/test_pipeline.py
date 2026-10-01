import pytest

from excel_copilot.llm.base import LLMError
from excel_copilot.llm.heuristic import HeuristicBackend
from excel_copilot.pipeline import Options, run_question

from .conftest import ScriptedBackend, plan, program

GOOD = program([{"op": "load", "out": "o", "table": "orders"},
                {"op": "aggregate", "out": "g", "input": "o", "group_by": ["Region"],
                 "aggregations": [{"column": "Revenue", "func": "sum", "as_name": "Revenue"}]}], "g",
               chart={"kind": "bar", "x": "Region", "y": "Revenue"})
BAD = program([{"op": "load", "out": "o", "table": "orders"},
               {"op": "aggregate", "out": "g", "input": "o", "group_by": ["Regoin"],
                "aggregations": [{"column": "Revenue", "func": "sum", "as_name": "Revenue"}]}], "g")


def events(backend, demo, q="revenue by region", storage=None, **opt):
    return list(run_question(q, demo.profile, demo.frames, backend, storage, options=Options(**opt)))


def types(evs):
    return [e.type for e in evs]


def test_happy_path_event_order(demo):
    evs = events(ScriptedBackend(plan(), [GOOD]), demo)
    assert types(evs) == ["stage", "context", "stage", "plan", "stage", "ir", "diagnostics", "code",
                          "stage", "log", "result", "done"]
    res = next(e for e in evs if e.type == "result").data
    assert res["row_count"] == 4 and len(res["rows"]) == 4 and [c["name"] for c in res["columns"]] == ["Region", "Revenue"]
    done = evs[-1].data
    assert done["ok"] and done["tokens"] == {"input": 300, "output": 100}


def test_repair_loop_feeds_diagnostics_back(demo):
    be = ScriptedBackend(plan(), [BAD, GOOD])
    evs = events(be, demo)
    diags = [e.data for e in evs if e.type == "diagnostics"]
    assert [d["ok"] for d in diags] == [False, True]
    assert be.histories[0] == []
    repair = be.histories[1]
    assert [t.role for t in repair] == ["assistant", "user"]
    assert "Regoin" in repair[0].content and "did you mean 'Region'" in repair[1].content
    assert evs[-1].data["ok"]


def test_gives_up_after_max_repairs(demo):
    evs = events(ScriptedBackend(plan(), [BAD, BAD, BAD]), demo, max_repairs=2)
    err = next(e for e in evs if e.type == "error").data
    assert err["stage"] == "validate" and "after 2 repair" in err["message"]
    assert evs[-1].data == {"ok": False, "run_id": evs[0].data["run_id"]}


def test_unanswerable_plan_stops_before_coding(demo):
    be = ScriptedBackend(plan(answerable=False, reason="no weather data"), [])
    evs = events(be, demo, q="weather?")
    assert "ir" not in types(evs) and be.histories == []
    assert evs[-1].data["answerable"] is False and evs[-1].data["reason"] == "no weather data"


def test_llm_errors_become_error_events(demo):
    evs = events(ScriptedBackend(LLMError("Claude API error 529: overloaded"), []), demo)
    assert [e.data["stage"] for e in evs if e.type == "error"] == ["plan"]
    assert not evs[-1].data["ok"]


def test_runtime_failure_triggers_one_repair(demo):
    boom = program([{"op": "load", "out": "a", "table": "orders"}, {"op": "load", "out": "b", "table": "orders"},
                    {"op": "join", "out": "j", "left": "a", "right": "b", "left_on": ["Region"],
                     "right_on": ["Region"], "how": "inner"}], "j")
    be = ScriptedBackend(plan(), [boom, GOOD])
    big = dict(demo.frames)
    big["orders"] = demo.frames["orders"].sample(n=6000, replace=True, random_state=0)
    evs = list(run_question("q", demo.profile, big, be))
    logs = [e.data["message"] for e in evs if e.type == "log"]
    assert any("join would produce" in m for m in logs)
    assert "runtime error: ValueError: join would produce" in be.histories[1][1].content
    assert evs[-1].data["ok"]


def test_artifacts_written(demo, tmp_path):
    from excel_copilot.storage import LocalStorage
    import openpyxl
    st = LocalStorage(tmp_path, "http://x", b"k")
    evs = events(ScriptedBackend(plan(), [GOOD]), demo, storage=st)
    arts = {e.data["name"]: e.data["key"] for e in evs if e.type == "artifact"}
    assert set(arts) == {"result.xlsx", "result.csv", "program.py", "run.json"}
    wb = openpyxl.load_workbook(tmp_path / arts["result.xlsx"])
    assert wb.sheetnames == ["Result", "Provenance"]
    assert [c.value for c in wb["Result"][4]] == ["Region", "Revenue"]
    prov = {r[0].value: r[1].value for r in wb["Provenance"].iter_rows()}
    assert prov["Question"] == "revenue by region" and "def run(tables):" in prov["Executed code"]
    assert (tmp_path / arts["program.py"]).read_text().startswith("def run(tables):")


@pytest.mark.parametrize("q,check", [
    ("total revenue by region", lambda df, raw: sorted(df["Revenue"]) == pytest.approx(sorted(
        sum(r["Revenue"] for r in raw if r["Region"] == g) for g in {r["Region"] for r in raw}))),
    ("how many orders were returned", lambda df, raw: df.iloc[0, 0] == sum(r["Status"] == "Returned" for r in raw)),
    ("average discount for West", lambda df, raw: df.iloc[0, 0] == pytest.approx(
        sum(r["Discount"] for r in raw if r["Region"] == "West") / sum(r["Region"] == "West" for r in raw))),
    ("top 3 SKUs by units in 2024", lambda df, raw: list(df["SKU"]) == [k for k, _ in sorted(
        {s: sum(int(r["Units"]) for r in raw if r["SKU"] == s) for s in {r["SKU"] for r in raw}}.items(),
        key=lambda kv: -kv[1])[:3]]),
])
def test_offline_backend_answers(demo, raw_orders, q, check):
    import pandas as pd
    evs = events(HeuristicBackend(), demo, q=q)
    res = next(e for e in evs if e.type == "result").data
    df = pd.DataFrame(res["rows"], columns=[c["name"] for c in res["columns"]])
    assert check(df, raw_orders), df


def test_offline_backend_declines_unrelated_questions(demo):
    evs = events(HeuristicBackend(), demo, q="what is the capital of France")
    assert evs[-1].data.get("answerable") is False
