from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pytest

from excel_copilot.ir.models import Plan, Program
from excel_copilot.llm.base import LLMError, Turn, Usage
from excel_copilot.profiler import profile_workbook

DEMO = Path(__file__).resolve().parents[1] / "examples" / "sales_ops.xlsx"


@pytest.fixture(scope="session")
def demo_bytes() -> bytes:
    return DEMO.read_bytes()


@pytest.fixture(scope="session")
def demo(demo_bytes):
    return profile_workbook(demo_bytes, "sales_ops.xlsx")


@pytest.fixture(scope="session")
def raw_orders(demo_bytes) -> list[dict]:
    """Orders read straight from the cells with openpyxl -- independent of the profiler."""
    ws = openpyxl.load_workbook(io.BytesIO(demo_bytes), data_only=True)["Orders"]
    header = [c.value for c in ws[3]]
    return [dict(zip(header, [c.value for c in row])) for row in ws.iter_rows(min_row=4, max_row=603)]


@pytest.fixture(scope="session")
def raw_sheet(demo_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(demo_bytes), data_only=True)

    def rows(sheet: str, header_row: int = 1) -> list[dict]:
        ws = wb[sheet]
        header = [c.value for c in ws[header_row]]
        return [dict(zip(header, [c.value for c in r])) for r in ws.iter_rows(min_row=header_row + 1)
                if any(c.value is not None for c in r)]
    return rows


def make_xlsx(build) -> bytes:
    wb = openpyxl.Workbook()
    build(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ScriptedBackend:
    """Replays canned plans/programs; records the repair history it was shown."""
    name = "scripted"

    def __init__(self, plan: Plan | Exception, programs: list[Program | Exception]):
        self._plan = plan
        self._programs = list(programs)
        self.histories: list[list[Turn]] = []

    def plan(self, question, context, profile):
        if isinstance(self._plan, Exception):
            raise self._plan
        return self._plan, Usage(model="scripted", input_tokens=100, output_tokens=20)

    def code(self, question, plan, context, profile, history):
        self.histories.append(list(history))
        if not self._programs:
            raise LLMError("script exhausted")
        p = self._programs.pop(0)
        if isinstance(p, Exception):
            raise p
        return p, Usage(model="scripted", input_tokens=200, output_tokens=80)


def program(steps: list[dict], output: str, chart=None, title="t") -> Program:
    return Program.model_validate({"steps": steps, "output": output, "title": title,
                                   "chart": chart or {"kind": "none", "x": "", "y": ""}})


def plan(tables=("orders",), answerable=True, reason="") -> Plan:
    return Plan(answerable=answerable, reason=reason, tables=list(tables), columns=[], steps=["x"], assumptions=[])
