"""Profile a workbook end to end: safe open -> regions -> columns -> formulas -> relationships.

`profile_workbook` returns both the JSON-serialisable profile and the typed
DataFrames for every table, so the executor never has to re-guess types.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import re
import time
import zipfile
from collections import Counter
from dataclasses import dataclass
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

from .columns import profile_column, to_number
from .formulas import analyse_column, readable
from .models import Issue, NamedRange, SheetProfile, TableProfile, WorkbookProfile
from .relationships import infer_relationships
from .tables import find_regions

MAX_UNCOMPRESSED = 300 * 1024 * 1024
MAX_RATIO = 250
MAX_ENTRIES = 5000
MAX_CELLS = 3_000_000


class UnsafeWorkbook(ValueError):
    """Raised for files that are not xlsx, look like zip bombs, or exceed limits."""


@dataclass
class ProfileResult:
    profile: WorkbookProfile
    frames: dict[str, pd.DataFrame]


def check_container(data: bytes) -> tuple[bool, list[Issue]]:
    """Validate the OOXML zip before openpyxl touches it. Returns (has_macros, issues)."""
    if not data.startswith(b"PK\x03\x04"):
        raise UnsafeWorkbook("not an .xlsx/.xlsm file (missing zip signature); legacy .xls is not supported")
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as e:
        raise UnsafeWorkbook(f"corrupt zip container: {e}") from e
    infos = zf.infolist()
    if len(infos) > MAX_ENTRIES:
        raise UnsafeWorkbook(f"too many zip entries ({len(infos)})")
    total = 0
    for i in infos:
        total += i.file_size
        if i.compress_size and i.file_size / i.compress_size > MAX_RATIO and i.file_size > 10_000_000:
            raise UnsafeWorkbook(f"suspicious compression ratio in {i.filename} (zip bomb?)")
        if ".." in i.filename or i.filename.startswith("/"):
            raise UnsafeWorkbook(f"illegal path in container: {i.filename}")
    if total > MAX_UNCOMPRESSED:
        raise UnsafeWorkbook(f"workbook expands to {total / 1e6:.0f} MB (limit {MAX_UNCOMPRESSED / 1e6:.0f} MB)")
    names = {i.filename for i in infos}
    if "xl/workbook.xml" not in names:
        raise UnsafeWorkbook("zip is not a spreadsheet (no xl/workbook.xml)")
    issues = []
    has_macros = any(n.endswith("vbaProject.bin") for n in names)
    if has_macros:
        issues.append(Issue(code="macros_present", severity="warning",
                            message="workbook contains VBA macros; they are never executed"))
    if any(n.startswith("xl/externalLinks/") for n in names):
        issues.append(Issue(code="external_links", severity="warning",
                            message="workbook links to external files; linked values may be stale"))
    return has_macros, issues


def _slug(s: str) -> str:
    s = re.sub(r"^tbl_?", "", s.strip(), flags=re.I) or s
    s = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", "_", s)
    s = re.sub(r"[^0-9a-zA-Z]+", "_", s).strip("_").lower()
    if not s or s[0].isdigit():
        s = "t_" + s
    return s


def _clean_headers(raw: list[Any], letters: list[str]) -> list[str]:
    out, seen = [], Counter()
    for h, letter in zip(raw, letters):
        name = re.sub(r"\s+", " ", str(h)).strip() if h is not None else ""
        if isinstance(h, (dt.date, dt.datetime)):
            name = h.strftime("%Y-%m")
        if not name:
            name = f"column_{letter}"
        seen[name.lower()] += 1
        if seen[name.lower()] > 1:
            name = f"{name}_{seen[name.lower()]}"
        out.append(name)
    return out


def coerce(values: list[Any], ctype: str) -> pd.Series:
    """Turn raw cell values into a typed pandas Series according to the profiled type."""
    if ctype in ("integer", "number", "currency", "percent"):
        nums = [to_number(v) if v is not None and not (isinstance(v, str) and v.strip().startswith("#")) else None
                for v in values]
        s = pd.Series(nums, dtype="Float64")
        if ctype == "integer":
            ok = s.dropna()
            if (ok == ok.round()).all():
                return s.astype("Int64")
        return s.astype("float64")
    if ctype in ("date", "datetime"):
        return pd.to_datetime(pd.Series([v if isinstance(v, (dt.date, dt.datetime)) else None
                                         for v in values], dtype="object"), errors="coerce")
    if ctype == "boolean":
        return pd.Series([v if isinstance(v, bool) else None for v in values], dtype="boolean")
    return pd.Series([None if v is None or (isinstance(v, str) and not v.strip()) else str(v).strip()
                      for v in values], dtype="string")


def profile_workbook(data: bytes, file_name: str = "workbook.xlsx") -> ProfileResult:
    t0 = time.perf_counter()
    has_macros, wb_issues = check_container(data)
    wb_v = load_workbook(io.BytesIO(data), data_only=True, keep_vba=False, keep_links=False)
    wb_f = load_workbook(io.BytesIO(data), data_only=False, keep_vba=False, keep_links=False)

    sheets: list[SheetProfile] = []
    tables: list[TableProfile] = []
    frames: dict[str, pd.DataFrame] = {}
    ids: Counter[str] = Counter()
    total_cells = 0
    missing_cache = 0
    lookups: list[tuple[str, list[str], str, tuple[str, str, str]]] = []

    for ws in wb_v.worksheets:
        wf = wb_f[ws.title]
        cells: dict[tuple[int, int], Any] = {}
        fmts: dict[tuple[int, int], str] = {}
        n_formula = n_error = 0
        for row in ws.iter_rows():
            for c in row:
                if c.value is None or (isinstance(c.value, str) and not c.value.strip()):
                    continue
                cells[(c.row, c.column)] = c.value
                fmts[(c.row, c.column)] = c.number_format
                if isinstance(c.value, str) and c.value.startswith("#") and c.value.endswith(("!", "?", "A")):
                    n_error += 1
        formulas: dict[tuple[int, int], str] = {}
        for row in wf.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas[(c.row, c.column)] = c.value
                    n_formula += 1
                    if (c.row, c.column) not in cells:
                        missing_cache += 1
                        cells[(c.row, c.column)] = None  # placeholder so regions include it
                elif hasattr(c.value, "text") and str(getattr(c.value, "text", "")).startswith("="):
                    formulas[(c.row, c.column)] = c.value.text  # array formulas
                    n_formula += 1
        total_cells += len(cells)
        if total_cells > MAX_CELLS:
            raise UnsafeWorkbook(f"workbook has more than {MAX_CELLS:,} non-empty cells")

        excel_tables = [(t.displayName or t.name, t.ref) for t in ws.tables.values()]
        regions, loose = find_regions(ws.title, {k: v for k, v in cells.items()}, excel_tables)
        merged = [str(m) for m in ws.merged_cells.ranges]

        sheet_table_ids = []
        for reg in regions:
            tid = _slug(reg.name)
            ids[tid] += 1
            if ids[tid] > 1:
                tid = f"{tid}_{ids[tid]}"
            letters = [get_column_letter(c) for c in range(reg.min_col, reg.max_col + 1)]
            raw_headers = ([cells.get((reg.header_row, c)) for c in range(reg.min_col, reg.max_col + 1)]
                           if reg.header_row else [None] * len(letters))
            names = _clean_headers(raw_headers, letters)
            columns = []
            series = {}
            t_issues: list[Issue] = []
            if reg.header_row and any(h is None for h in raw_headers):
                t_issues.append(Issue(code="blank_header", severity="warning",
                                      message="some columns have no header; named column_<letter>"))
            dupes = [n for n in names if re.search(r"_\d+$", n) and n.rsplit("_", 1)[0] in names]
            if dupes:
                t_issues.append(Issue(code="duplicate_header", severity="warning",
                                      message=f"duplicate headers renamed: {', '.join(dupes)}"))
            for i, (col, letter, name, raw) in enumerate(zip(range(reg.min_col, reg.max_col + 1),
                                                             letters, names, raw_headers)):
                vals = [cells.get((r, col)) for r in reg.data_rows]
                cp = profile_column(name=name, header=None if raw is None else str(raw), letter=letter,
                                    index=i, sheet=ws.title, rows=reg.data_rows, values=vals,
                                    formats=[fmts.get((r, col)) for r in reg.data_rows])
                fa = analyse_column([(r, formulas.get((r, col))) for r in reg.data_rows], len(reg.data_rows))
                if fa.formula:
                    cp.formula, cp.formula_coverage = fa.formula, round(fa.coverage, 4)
                    cp.references = fa.cross_sheet_refs
                    if fa.lookup:
                        lookups.append((tid, letters, name, fa.lookup))
                    if 0 < fa.coverage < 0.95:
                        cp.issues.append(Issue(code="inconsistent_formula", severity="warning",
                                               message=f"formula covers only {fa.coverage:.0%} of rows; "
                                                       "other rows are hard-coded or use a different formula"))
                columns.append(cp)
                series[name] = coerce(vals, cp.type)
            letter_to_name = {c.letter: c.name for c in columns}
            for cp in columns:
                if cp.formula:
                    cp.formula_readable = readable(cp.formula, letter_to_name)
            tp = TableProfile(id=tid, sheet=ws.title, name=reg.name, source=reg.source,  # type: ignore[arg-type]
                              range=reg.ref, header_row=reg.header_row or reg.min_row,
                              n_rows=len(reg.data_rows), columns=columns, title=reg.title,
                              layout=reg.layout, issues=t_issues)  # type: ignore[arg-type]
            if reg.layout == "wide":
                tp.issues.append(Issue(code="wide_format", severity="info",
                                       message="period columns (e.g. months) are laid out across; "
                                               "unpivot before comparing with long tables"))
            merged_inside = [m for m in merged if _overlaps(m, reg) and range_boundaries(m)[1] > reg.min_row]
            if merged_inside:
                tp.issues.append(Issue(code="merged_cells_in_table", severity="warning",
                                       message=f"merged cells inside the table: {', '.join(merged_inside[:5])}"))
            tables.append(tp)
            frames[tid] = pd.DataFrame(series)
            sheet_table_ids.append(tid)

        sheets.append(SheetProfile(
            name=ws.title, state=ws.sheet_state, dimensions=ws.dimensions,  # type: ignore[arg-type]
            max_row=ws.max_row, max_column=ws.max_column, tables=sheet_table_ids, merged_ranges=merged[:50],
            loose_cells=[(f"{get_column_letter(c)}{r}", _short(cells[(r, c)])) for r, c in loose[:40]
                         if cells.get((r, c)) is not None],
            formula_cells=n_formula, error_cells=n_error))

    if missing_cache:
        wb_issues.append(Issue(code="no_cached_values", severity="warning", count=missing_cache,
                               message=f"{missing_cache} formula cell(s) have no cached value (file saved "
                                       "by a library, not Excel); their values read as empty"))
    for s in sheets:
        if s.state != "visible":
            wb_issues.append(Issue(code="hidden_sheet", severity="info", location=s.name,
                                   message=f"sheet '{s.name}' is {s.state}"))

    named = []
    for name, dn in wb_f.defined_names.items():
        value = None
        try:
            dests = list(dn.destinations)
            if len(dests) == 1:
                sh, ref = dests[0]
                c0, r0, c1, r1 = range_boundaries(ref.replace("$", ""))
                if (c0, r0) == (c1, r1) and sh in wb_v.sheetnames:
                    value = _short(wb_v[sh].cell(r0, c0).value)
        except Exception:
            pass
        named.append(NamedRange(name=name, refers_to=dn.attr_text, value=value))

    rels = infer_relationships(tables, frames, lookups)
    profile = WorkbookProfile(
        file_name=file_name, sha256=hashlib.sha256(data).hexdigest(), size_bytes=len(data),
        sheets=sheets, tables=tables, relationships=rels, named_ranges=named, issues=wb_issues,
        has_macros=has_macros, profiled_ms=round((time.perf_counter() - t0) * 1000, 1))
    return ProfileResult(profile, frames)


def _overlaps(merged: str, reg) -> bool:
    c0, r0, c1, r1 = range_boundaries(merged)
    return not (c1 < reg.min_col or c0 > reg.max_col or r1 < reg.min_row or r0 > reg.max_row)


def _short(v: Any) -> str:
    if isinstance(v, (dt.date, dt.datetime)):
        return v.isoformat()
    s = str(v)
    return s if len(s) <= 60 else s[:57] + "..."
