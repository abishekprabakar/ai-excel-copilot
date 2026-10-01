"""Find the tables inside a worksheet.

Real workbooks rarely start at A1 with a clean header. We take Excel Tables
(ListObjects) as ground truth, then find the remaining data regions as
8-connected components of non-empty cells, and locate each region's header row
(skipping title rows glued on top of it). Single cells and single rows left over
become "loose cells" (notes, parameters like an FX rate) rather than tables.
"""
from __future__ import annotations

import re
from collections import deque
from dataclasses import dataclass, field
from typing import Any

from openpyxl.utils import get_column_letter, range_boundaries

_MONTHS = {m.lower() for m in ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep",
                                 "Oct", "Nov", "Dec", "January", "February", "March", "April",
                                 "June", "July", "August", "September", "October", "November",
                                 "December", "Sept"]}
_PERIOD = re.compile(r"^(?:(?:19|20)\d{2}|q[1-4](?:[\s-]*(?:19|20)?\d{2})?|(?:fy|cy)\s*'?\d{2,4}|"
                     r"[a-z]{3,9}[\s-]*'?\d{2,4})$", re.I)


@dataclass
class Region:
    sheet: str
    min_row: int
    min_col: int
    max_row: int
    max_col: int
    header_row: int | None
    source: str                      # "excel_table" | "detected"
    name: str
    title: str | None = None
    layout: str = "long"
    data_rows: list[int] = field(default_factory=list)

    @property
    def ref(self) -> str:
        return (f"{get_column_letter(self.min_col)}{self.min_row}:"
                f"{get_column_letter(self.max_col)}{self.max_row}")


def is_period_label(v: Any) -> bool:
    if isinstance(v, int) and 1900 <= v <= 2100:
        return True
    if not isinstance(v, str):
        return False
    s = v.strip().lower()
    return s in _MONTHS or bool(_PERIOD.match(s))


def detect_layout(headers: list[Any]) -> str:
    """'wide' if >= 4 consecutive headers are periods (months, years, quarters)."""
    run = best = 0
    for h in headers:
        run = run + 1 if is_period_label(h) else 0
        best = max(best, run)
    return "wide" if best >= 4 else "long"


def _is_header_value(v: Any) -> bool:
    return isinstance(v, str) and bool(v.strip()) or is_period_label(v)


def _components(cells: set[tuple[int, int]]) -> list[set[tuple[int, int]]]:
    seen: set[tuple[int, int]] = set()
    comps = []
    for start in sorted(cells):
        if start in seen:
            continue
        comp = {start}
        seen.add(start)
        q = deque([start])
        while q:
            r, c = q.popleft()
            for dr in (-1, 0, 1):
                for dc in (-1, 0, 1):
                    nb = (r + dr, c + dc)
                    if nb in cells and nb not in seen:
                        seen.add(nb)
                        comp.add(nb)
                        q.append(nb)
        comps.append(comp)
    return comps


def find_regions(sheet: str, cells: dict[tuple[int, int], Any],
                 excel_tables: list[tuple[str, str]]) -> tuple[list[Region], list[tuple[int, int]]]:
    """Return (table regions, coordinates of loose cells). `cells` holds non-empty values."""
    regions: list[Region] = []
    claimed: set[tuple[int, int]] = set()
    for name, ref in excel_tables:
        c0, r0, c1, r1 = range_boundaries(ref)
        headers = [cells.get((r0, c)) for c in range(c0, c1 + 1)]
        regions.append(Region(sheet, r0, c0, r1, c1, r0, "excel_table", name,
                              layout=detect_layout(headers),
                              data_rows=list(range(r0 + 1, r1 + 1))))
        claimed |= {(r, c) for r in range(r0, r1 + 1) for c in range(c0, c1 + 1)}

    loose: list[tuple[int, int]] = []
    for comp in _components(set(cells) - claimed):
        rows = sorted({r for r, _ in comp})
        cols = sorted({c for _, c in comp})
        r0, r1, c0, c1 = rows[0], rows[-1], cols[0], cols[-1]
        width = c1 - c0 + 1
        header = None
        for r in range(r0, min(r0 + 6, r1)):
            vals = [cells.get((r, c)) for c in range(c0, c1 + 1)]
            filled = [v for v in vals if v is not None]
            if len(filled) >= max(2 if width > 1 else 1, 0.5 * width) and all(map(_is_header_value, filled)):
                header = r
                break
        if header is None:
            # no header: a block of >= 3 rows and >= 2 columns is still a table
            if len(rows) >= 3 and width >= 2:
                header = r0 - 1          # virtual header above the block
            else:
                loose.extend(sorted(comp))
                continue
        data_rows = [r for r in range(header + 1, r1 + 1)
                     if any((r, c) in cells for c in range(c0, c1 + 1))]
        if not data_rows or (width == 1 and len(data_rows) < 2):
            loose.extend(sorted(comp))
            continue
        title_cells = [str(cells[(r, c)]) for r in range(r0, header) for c in range(c0, c1 + 1)
                       if (r, c) in cells]
        headers = [cells.get((header, c)) for c in range(c0, c1 + 1)]
        regions.append(Region(sheet, max(header, r0), c0, r1, c1,
                              header if header >= r0 else None, "detected", sheet,
                              title=" ".join(title_cells) or None, layout=detect_layout(headers),
                              data_rows=data_rows))

    # A lone text cell directly above a region (e.g. a merged title row separated by a
    # blank line) is that region's title.
    for reg in regions:
        if reg.title is None and reg.source == "detected":
            for r in range(reg.min_row - 1, max(reg.min_row - 3, 0), -1):
                hits = [(r, c) for (rr, c) in loose if rr == r and reg.min_col <= c <= reg.max_col]
                if len(hits) == 1 and isinstance(cells[hits[0]], str):
                    reg.title = cells[hits[0]]
                    loose.remove(hits[0])
                    break
    return regions, loose
