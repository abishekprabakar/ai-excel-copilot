"""The question-answering pipeline as a stream of events.

    context -> plan (LLM) -> code (LLM) -> validate+compile -> [repair loop] -> verify
            -> execute (sandbox) -> artifacts -> done

`run_question` is a generator of `Event`s. The API serialises them as
server-sent events; the CLI prints them. Nothing is buffered until the end, so
a client sees each stage as it happens.
"""
from __future__ import annotations

import datetime as dt
import io
import json
import time
import uuid
from dataclasses import asdict, dataclass, field
from typing import Any, Iterator

import pandas as pd

from .graph import render_context
from .ir.compiler import compile_program
from .ir.models import Plan, Program
from .llm.base import Backend, LLMError, Turn, Usage
from .llm.prompts import repair_user
from .profiler.models import WorkbookProfile
from .sandbox.executor import Limits, execute
from .sandbox.verify import VerificationError, verify
from .storage import Storage


@dataclass
class Event:
    type: str                  # stage | log | context | plan | ir | diagnostics | code | result | artifact | error | done
    data: dict[str, Any]
    ts: float = field(default_factory=time.time)

    def sse(self) -> str:
        return f"event: {self.type}\ndata: {json.dumps(self.data, default=str)}\n\n"


@dataclass
class Options:
    context_budget: int = 1500
    coder_budget: int = 2500
    max_repairs: int = 2
    preview_rows: int = 200
    limits: Limits = field(default_factory=Limits)


def _usage(u: Usage) -> dict:
    return asdict(u)


def _preview(df: pd.DataFrame, n: int) -> dict:
    head = df.head(n)
    cols = [{"name": str(c), "dtype": str(t)} for c, t in head.dtypes.items()]
    rows = json.loads(head.to_json(orient="values", date_format="iso"))
    return {"columns": cols, "rows": rows}


def result_workbook(df: pd.DataFrame, *, question: str, title: str, plan: Plan, program: Program,
                    source: str, profile: WorkbookProfile) -> bytes:
    """An .xlsx with the result and a Provenance sheet describing how it was produced."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Result"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Question: {question}"])
    ws.append([])
    header_row = 4
    ws.append([str(c) for c in df.columns])
    for c in ws[header_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for rec in df.itertuples(index=False):
        row = []
        for v in rec:
            if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NA or v is pd.NaT:
                row.append(None)
            elif isinstance(v, pd.Timestamp):
                row.append(v.to_pydatetime().replace(tzinfo=None))
            elif hasattr(v, "item"):
                row.append(v.item())
            else:
                row.append(v)
        ws.append(row)
    for j, (name, dtype) in enumerate(df.dtypes.items(), start=1):
        letter = get_column_letter(j)
        fmt = None
        if pd.api.types.is_datetime64_any_dtype(dtype):
            fmt = "yyyy-mm-dd"
        elif pd.api.types.is_float_dtype(dtype):
            lname = str(name).lower()
            fmt = "0.0%" if ("%" in lname or "margin" in lname or "rate" in lname) and df[name].abs().max(skipna=True) <= 5 \
                else "#,##0.00"
        elif pd.api.types.is_integer_dtype(dtype):
            fmt = "#,##0"
        if fmt:
            for cell in ws[letter][header_row:]:
                cell.number_format = fmt
        width = max([len(str(name))] + [len(str(v)) for v in df[name].head(200).tolist()]) + 2
        ws.column_dimensions[letter].width = min(max(width, 8), 50)
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(max(1, len(df.columns)))}{header_row + len(df)}"

    pv = wb.create_sheet("Provenance")
    rows = [("Question", question), ("Source workbook", profile.file_name), ("Source SHA-256", profile.sha256),
            ("Generated (UTC)", dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d %H:%M:%S")),
            ("Plan", "\n".join(f"{i + 1}. {s}" for i, s in enumerate(plan.steps))),
            ("Assumptions", "\n".join(plan.assumptions) or "-"),
            ("IR program", program.model_dump_json(indent=1)), ("Executed code", source)]
    for k, v in rows:
        pv.append([k, v])
    pv.column_dimensions["A"].width = 18
    pv.column_dimensions["B"].width = 110
    for r in pv.iter_rows():
        r[0].font = Font(bold=True)
        r[1].alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run_question(question: str, profile: WorkbookProfile, frames: dict[str, pd.DataFrame],
                 backend: Backend, storage: Storage | None = None, *, workbook_id: str = "local",
                 options: Options | None = None, run_id: str | None = None) -> Iterator[Event]:
    opt = options or Options()
    run_id = run_id or uuid.uuid4().hex
    t_start = time.perf_counter()
    trace: dict[str, Any] = {"run_id": run_id, "workbook_id": workbook_id, "question": question,
                             "backend": backend.name, "usage": []}
    yield Event("stage", {"stage": "start", "run_id": run_id, "backend": backend.name})

    def fail(stage: str, message: str, **extra) -> Event:
        trace["error"] = {"stage": stage, "message": message}
        return Event("error", {"stage": stage, "message": message, **extra})

    # 1. context
    ctx = render_context(profile, question, opt.context_budget)
    yield Event("context", {"tokens": ctx.tokens, "levels": ctx.levels, "text": ctx.text})

    # 2. plan
    yield Event("stage", {"stage": "plan", "status": "running"})
    try:
        plan, usage = backend.plan(question, ctx.text, profile)
    except LLMError as e:
        yield fail("plan", str(e))
        yield Event("done", {"ok": False, "run_id": run_id})
        return
    trace["plan"] = plan.model_dump()
    trace["usage"].append({"stage": "plan", **_usage(usage)})
    yield Event("plan", {"plan": plan.model_dump(), "usage": _usage(usage)})
    if not plan.answerable:
        yield Event("done", {"ok": False, "run_id": run_id, "answerable": False, "reason": plan.reason})
        return

    # 3. code -> validate -> repair. The coder sees full detail for the planned tables.
    known = [t.id for t in profile.tables]
    focus = [t for t in plan.tables if t in known] or None
    coder_ctx = render_context(profile, question, opt.coder_budget, only_tables=focus, min_level=1) if focus else ctx
    history: list[Turn] = []
    compiled = program = None
    for attempt in range(opt.max_repairs + 1):
        yield Event("stage", {"stage": "code", "status": "running", "attempt": attempt + 1})
        try:
            program, usage = backend.code(question, plan, coder_ctx.text, profile, history)
        except LLMError as e:
            yield fail("code", str(e))
            yield Event("done", {"ok": False, "run_id": run_id})
            return
        trace["usage"].append({"stage": f"code#{attempt + 1}", **_usage(usage)})
        yield Event("ir", {"attempt": attempt + 1, "program": program.model_dump(), "usage": _usage(usage)})
        compiled = compile_program(program, profile)
        yield Event("diagnostics", {"attempt": attempt + 1, "ok": compiled.ok,
                                    "items": [d.model_dump() for d in compiled.diagnostics]})
        if compiled.ok:
            break
        history += [Turn("assistant", program.model_dump_json()),
                    Turn("user", repair_user([d.render() for d in compiled.errors]))]
    assert program is not None and compiled is not None
    trace["program"] = program.model_dump()
    if not compiled.ok:
        yield fail("validate", f"program still invalid after {opt.max_repairs} repair(s)",
                   items=[d.model_dump() for d in compiled.errors])
        yield Event("done", {"ok": False, "run_id": run_id})
        return

    # 4. verify + execute (one more repair round if execution itself fails)
    for exec_attempt in range(2):
        try:
            verify(compiled.source)
        except VerificationError as e:
            yield fail("verify", str(e))
            yield Event("done", {"ok": False, "run_id": run_id})
            return
        trace["code"] = compiled.source
        yield Event("code", {"source": compiled.source, "tables": compiled.tables_used})
        yield Event("stage", {"stage": "execute", "status": "running"})
        res = execute(compiled.source, {t: frames[t] for t in compiled.tables_used}, opt.limits)
        yield Event("log", {"stage": "execute", "message": f"sandbox finished in {res.wall_ms:.0f} ms "
                                                           f"(program {res.exec_ms:.0f} ms)"})
        if res.ok:
            break
        if exec_attempt == 0 and opt.max_repairs > 0:
            yield Event("log", {"stage": "execute", "message": f"execution failed: {res.error}; asking the coder to repair"})
            history += [Turn("assistant", program.model_dump_json()),
                        Turn("user", repair_user([f"runtime error: {res.error}"]))]
            try:
                program, usage = backend.code(question, plan, coder_ctx.text, profile, history)
            except LLMError as e:
                yield fail("execute", f"{res.error} (repair failed: {e})")
                yield Event("done", {"ok": False, "run_id": run_id})
                return
            trace["usage"].append({"stage": "code#runtime-repair", **_usage(usage)})
            yield Event("ir", {"attempt": "runtime-repair", "program": program.model_dump(), "usage": _usage(usage)})
            compiled = compile_program(program, profile)
            yield Event("diagnostics", {"attempt": "runtime-repair", "ok": compiled.ok,
                                        "items": [d.model_dump() for d in compiled.diagnostics]})
            if not compiled.ok:
                break
            continue
        break
    if not compiled.ok or not res.ok or res.frame is None:
        yield fail("execute", res.error or "repaired program was invalid", stderr=res.stderr[-1000:])
        yield Event("done", {"ok": False, "run_id": run_id})
        return

    df = res.frame
    trace["result"] = {"rows": res.rows, "truncated": res.truncated, "columns": [str(c) for c in df.columns]}
    yield Event("result", {"title": program.title, "row_count": res.rows, "truncated": res.truncated,
                           "chart": program.chart.model_dump(), **_preview(df, opt.preview_rows)})

    # 5. artifacts
    artifacts = []
    if storage is not None:
        files = {
            "result.xlsx": (result_workbook(df, question=question, title=program.title, plan=plan, program=program,
                                            source=compiled.source, profile=profile),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            "result.csv": (df.to_csv(index=False).encode(), "text/csv"),
            "program.py": (compiled.source.encode(), "text/x-python"),
        }
        trace["elapsed_ms"] = round((time.perf_counter() - t_start) * 1000, 1)
        files["run.json"] = (json.dumps(trace, indent=1, default=str).encode(), "application/json")
        for name, (data, ctype) in files.items():
            key = f"runs/{run_id}/{name}"
            storage.put(key, data, ctype)
            art = {"name": name, "key": key, "bytes": len(data), "content_type": ctype}
            artifacts.append(art)
            yield Event("artifact", art)

    total_in = sum(u.get("input_tokens", 0) for u in trace["usage"])
    total_out = sum(u.get("output_tokens", 0) for u in trace["usage"])
    yield Event("done", {"ok": True, "run_id": run_id, "rows": res.rows,
                         "elapsed_ms": round((time.perf_counter() - t_start) * 1000, 1),
                         "tokens": {"input": total_in, "output": total_out}, "artifacts": [a["name"] for a in artifacts]})
