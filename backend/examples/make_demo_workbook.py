"""Generate examples/sales_ops.xlsx: a deliberately realistic, slightly messy workbook.

What makes it realistic (and what the profiler is expected to find):
  * Orders: a title row and a blank row above the header, 600 rows, filled-down
    formulas (Revenue, Unit Cost via cross-sheet VLOOKUP, Margin %), a #DIV/0!
    from a zero-unit order, numbers stored as text, two orphan customer IDs.
  * Customers: a real Excel Table (ListObject) with an Email column (PII).
  * Products: a plain lookup table referenced by Orders formulas.
  * Targets: a wide month-by-region matrix and a named range (FX_EUR_USD).
  * Lookups: a hidden sheet.

openpyxl writes formulas without cached values, so the file is recalculated with
LibreOffice (if installed) so it looks like a file saved by Excel. Deterministic.

    python examples/make_demo_workbook.py [--out examples/sales_ops.xlsx] [--no-recalc]
"""
from __future__ import annotations

import argparse
import datetime as dt
import random
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

REGIONS = ["West", "East", "North", "South"]
SEGMENTS = ["Enterprise", "SMB", "Consumer"]
COUNTRIES = ["US", "US", "US", "CA", "GB", "DE"]
CATEGORIES = {"Hardware": (40, 220), "Software": (5, 60), "Services": (80, 400)}
FIRST = ["Ava", "Liam", "Maya", "Noah", "Zoe", "Ethan", "Iris", "Omar", "Lena", "Ravi",
         "Sofia", "Kai", "Nora", "Leo", "Priya", "Owen"]
LAST = ["Patel", "Nguyen", "Smith", "Garcia", "Kim", "Okafor", "Rossi", "Chen", "Silva", "Novak"]


def build(path: Path) -> None:
    rng = random.Random(2024)
    wb = Workbook()

    # ------------------------------------------------------------- Products
    ws_p = wb.active
    ws_p.title = "Products"
    ws_p.append(["SKU", "Product Name", "Category", "Unit Cost", "List Price"])
    products = []
    for i in range(25):
        cat = list(CATEGORIES)[i % 3]
        lo, hi = CATEGORIES[cat]
        price = round(rng.uniform(lo, hi), 2)
        cost = round(price * rng.uniform(0.35, 0.7), 2)
        sku = f"SKU-{1000 + i}"
        products.append((sku, price))
        ws_p.append([sku, f"{cat} item {i + 1}", cat, cost, price])
    for row in ws_p.iter_rows(min_row=2, min_col=4, max_col=5):
        for c in row:
            c.number_format = '"$"#,##0.00'

    # ------------------------------------------------------------- Customers
    ws_c = wb.create_sheet("Customers")
    ws_c.append(["Customer ID", "Name", "Email", "Segment", "Country", "Signup Date"])
    for cid in range(1, 81):
        first, last = rng.choice(FIRST), rng.choice(LAST)
        signup = dt.date(2021, 1, 1) + dt.timedelta(days=rng.randint(0, 1000))
        ws_c.append([cid, f"{first} {last}", f"{first}.{last}{cid}@example.com".lower(),
                     rng.choice(SEGMENTS), rng.choice(COUNTRIES), signup])
        ws_c.cell(row=cid + 1, column=6).number_format = "yyyy-mm-dd"
    tbl = Table(displayName="tblCustomers", ref=f"A1:F{ws_c.max_row}")
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws_c.add_table(tbl)

    # ------------------------------------------------------------- Orders
    ws_o = wb.create_sheet("Orders", 0)
    ws_o["A1"] = "Sales Orders FY2024"
    ws_o["A1"].font = Font(bold=True, size=14)
    ws_o.merge_cells("A1:L1")
    header = ["Order ID", "Order Date", "Customer ID", "SKU", "Region", "Units", "Unit Price",
              "Discount", "Revenue", "Status", "Unit Cost", "Margin %"]
    ws_o.append([])
    ws_o.append(header)
    for c in ws_o[3]:
        c.font = Font(bold=True)
    start = dt.date(2024, 1, 1)
    for n in range(600):
        r = n + 4
        sku, price = rng.choice(products)
        cid = rng.randint(1, 80)
        if n in (77, 412):
            cid = 90 + n % 7            # orphan customer IDs
        units = rng.randint(1, 40)
        if n == 250:
            units = 0                   # produces #DIV/0! in Margin %
        units_cell: object = units
        if n in (31, 32, 333):
            units_cell = str(units)     # numbers stored as text
        discount = rng.choice([0, 0, 0, 0.05, 0.1, 0.15])
        status = rng.choices(["Shipped", "Shipped", "Shipped", "Pending", "Returned"])[0]
        ws_o.append([
            10000 + n,
            start + dt.timedelta(days=rng.randint(0, 365)),
            cid, sku, rng.choice(REGIONS), units_cell,
            round(price * rng.uniform(0.95, 1.05), 2), discount,
            f"=F{r}*G{r}*(1-H{r})",
            status,
            f"=VLOOKUP(D{r},Products!$A$2:$E$26,4,FALSE)",
            f"=(I{r}-F{r}*K{r})/I{r}",
        ])
        ws_o.cell(row=r, column=2).number_format = "yyyy-mm-dd"
        for col in (7, 9, 11):
            ws_o.cell(row=r, column=col).number_format = '"$"#,##0.00'
        for col in (8, 12):
            ws_o.cell(row=r, column=col).number_format = "0.0%"

    # ------------------------------------------------------------- Targets (wide)
    ws_t = wb.create_sheet("Targets")
    months = [dt.date(2024, m, 1).strftime("%b") for m in range(1, 13)]
    ws_t.append(["Region"] + months)
    for reg in REGIONS:
        base = rng.randint(18000, 30000)
        ws_t.append([reg] + [round(base * (1 + 0.02 * m) / 1000) * 1000 for m in range(12)])
    ws_t["A8"] = "EUR/USD"
    ws_t["B8"] = 1.09
    wb.defined_names["FX_EUR_USD"] = DefinedName("FX_EUR_USD", attr_text="Targets!$B$8")

    # ------------------------------------------------------------- hidden lookups
    ws_l = wb.create_sheet("Lookups")
    ws_l.append(["Status", "Counts As Sale"])
    for s, v in [("Shipped", True), ("Pending", False), ("Returned", False)]:
        ws_l.append([s, v])
    ws_l.sheet_state = "hidden"

    wb.save(path)


def recalc_with_libreoffice(path: Path) -> bool:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return False
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                        "--outdir", tmp, str(path)], check=True, capture_output=True, timeout=180)
        shutil.copy(Path(tmp) / path.name, path)
    return True


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(Path(__file__).with_name("sales_ops.xlsx")))
    ap.add_argument("--no-recalc", action="store_true")
    args = ap.parse_args()
    out = Path(args.out)
    build(out)
    recalced = False if args.no_recalc else recalc_with_libreoffice(out)
    print(f"wrote {out} ({'recalculated with LibreOffice' if recalced else 'formulas without cached values'})")


if __name__ == "__main__":
    main()
